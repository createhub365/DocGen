"""
Walkthrough: auto-suggest is frontend-only; upload + GET must not persist.

Mirrors frontend mappingSuggestions.js (case-insensitive exact match).
Bootstraps the same in-memory SQLite isolation as tests/conftest.py
(does NOT touch docgen.db / live Postgres).

Run from docgen/backend:
  python scripts/verify_mapping_suggest.py
"""
from __future__ import annotations

import io
import os
import shutil
import sys
import tempfile
from pathlib import Path

# Isolate before any app imports — same contract as tests/conftest.py
os.environ["DATABASE_URL"] = "sqlite:///:memory:"
os.environ["JWT_SECRET"] = (
    "test_jwt_secret_for_platform_isolation_tests_only_64chars_xx"
)
os.environ["ENVIRONMENT"] = "development"
os.environ["ALLOW_DEMO_SEED"] = "false"
os.environ["DOCGEN_SKIP_PDF"] = "true"
os.environ["PLATFORM_SIGNUP_RATE_LIMIT"] = "1000/hour"
os.environ["PLATFORM_LOGIN_RATE_LIMIT"] = "1000/minute"
os.environ["PLATFORM_INVITE_RATE_LIMIT"] = "1000/hour"

_TEST_FILE_ROOT = Path(tempfile.mkdtemp(prefix="docgen_suggest_walk_"))
_TEST_TEMPLATE_DIR = _TEST_FILE_ROOT / "template_store"
_TEST_OUTPUT_DIR = _TEST_FILE_ROOT / "output"
_TEST_TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
_TEST_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
os.environ["TEMPLATE_DIR"] = str(_TEST_TEMPLATE_DIR)
os.environ["OUTPUT_DIR"] = str(_TEST_OUTPUT_DIR)

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from sqlalchemy import create_engine, event  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402
from sqlalchemy.pool import StaticPool  # noqa: E402

import database  # noqa: E402
from database import Base, get_db  # noqa: E402

_engine = create_engine(
    "sqlite://",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)


@event.listens_for(_engine, "connect")
def _fk_on(dbapi_connection, connection_record):
    cursor = dbapi_connection.cursor()
    cursor.execute("PRAGMA foreign_keys=ON")
    cursor.close()


database.engine = _engine
database.SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=_engine)

import models  # noqa: E402,F401
from auth import create_org_jwt, hash_password  # noqa: E402
from main import app  # noqa: E402
from models import Organization, OrgUser, User  # noqa: E402

SessionLocal = database.SessionLocal
Base.metadata.create_all(bind=_engine)


def suggest_field_key(placeholder_key: str, resolvable_keys: list[str]) -> str | None:
    needle = (placeholder_key or "").lower()
    if not needle:
        return None
    for key in resolvable_keys:
        if str(key).lower() == needle:
            return key
    return None


def build_initial_selections(detected, persisted, resolvable_keys):
    saved = {
        row["placeholder_key"]: row["field_key"]
        for row in persisted
        if row.get("is_mapped") and row.get("placeholder_key") and row.get("field_key")
    }
    selections = {}
    suggested = []
    for key in detected:
        if key in saved:
            selections[key] = saved[key]
            continue
        match = suggest_field_key(key, resolvable_keys)
        if match:
            selections[key] = match
            suggested.append(key)
        else:
            selections[key] = ""
    return selections, suggested


def _docx_bytes(*placeholders: str) -> bytes:
    doc = Document()
    doc.add_paragraph(" ".join(f"{{{{{p}}}}}" for p in placeholders))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_org_admin(db, *, slug: str, username: str, org_name: str):
    user = User(
        username=username,
        full_name=f"{org_name} Admin",
        password_hash=hash_password("test-password-123"),
        role="staff",
        is_active=True,
    )
    db.add(user)
    db.flush()
    org = Organization(name=org_name, slug=slug, is_active=True)
    db.add(org)
    db.flush()
    membership = OrgUser(org_id=org.id, user_id=user.id, role="org_admin")
    db.add(membership)
    db.commit()
    token = create_org_jwt(user_id=user.id, org_id=org.id, role="org_admin")
    return {"token": token, "org": org}


def main() -> int:
    db = SessionLocal()
    try:
        org = _make_org_admin(
            db,
            slug="suggest-walk",
            username="suggest_admin",
            org_name="Suggest Walk Org",
        )
    finally:
        db.close()

    def override_get_db():
        session = SessionLocal()
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app)
    client.cookies.set("platform_access_token", org["token"])

    try:
        # Pure logic
        assert suggest_field_key("Company_Name", ["company_name"]) == "company_name"
        assert suggest_field_key("Other_Thing", ["company_name"]) is None

        dt = client.post(
            "/api/platform/document-types/",
            json={"name": "Suggest Walk", "slug": "suggest-walk"},
        )
        assert dt.status_code == 201, dt.text
        dt_id = dt.json()["id"]

        flow = client.post(f"/api/platform/{dt_id}/flow", json={})
        assert flow.status_code == 201, flow.text
        flow_id = flow.json()["id"]

        step = client.post(
            f"/api/platform/{flow_id}/steps",
            json={
                "step_type": "text_field",
                "order_index": 0,
                "label": "Company",
                "is_enabled": True,
            },
        )
        assert step.status_code == 201, step.text
        step_id = step.json()["id"]

        field = client.post(
            f"/api/platform/steps/{step_id}/fields",
            json={
                "field_key": "company_name",
                "field_label": "Company name",
                "field_type": "text",
                "is_required": True,
            },
        )
        assert field.status_code == 201, field.text

        pub = client.post(f"/api/platform/{flow_id}/publish")
        assert pub.status_code == 200, pub.text

        upload = client.post(
            f"/api/platform/{dt_id}/templates",
            files={
                "file": (
                    "suggest.docx",
                    _docx_bytes("company_name", "Other_Thing"),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert upload.status_code == 201, upload.text
        template_id = upload.json()["id"]

        before = client.get(f"/api/platform/templates/{template_id}/mappings")
        assert before.status_code == 200, before.text
        body = before.json()
        detected = body["detected_placeholders"]
        assert "company_name" in detected
        assert "Other_Thing" in detected
        mapped = [m for m in body.get("mappings") or [] if m.get("is_mapped")]
        assert mapped == [], "GET after upload must not persist mappings"

        selections, suggested = build_initial_selections(
            detected, body.get("mappings") or [], ["company_name"]
        )
        assert selections["company_name"] == "company_name"
        assert "company_name" in suggested
        assert selections["Other_Thing"] == ""
        assert "Other_Thing" not in suggested

        after = client.get(f"/api/platform/templates/{template_id}/mappings")
        assert [
            m for m in after.json().get("mappings") or [] if m.get("is_mapped")
        ] == [], "still unsaved after UI-suggest simulation"

        save = client.post(
            f"/api/platform/templates/{template_id}/mappings",
            json={
                "mappings": [
                    {"placeholder_key": "company_name", "field_key": "company_name"},
                ]
            },
        )
        assert save.status_code == 200, save.text
        assert any(
            m.get("is_mapped") and m["placeholder_key"] == "company_name"
            for m in save.json().get("mappings") or []
        )

        print(
            {
                "matching": "case-insensitive exact",
                "suggested_visual": "blue AntD Tag 'Suggested'",
                "preselected": "company_name",
                "empty": "Other_Thing",
                "auto_post_on_load": False,
                "explicit_save": True,
                "ok": True,
            }
        )
        return 0
    finally:
        app.dependency_overrides.clear()
        shutil.rmtree(_TEST_FILE_ROOT, ignore_errors=True)


if __name__ == "__main__":
    raise SystemExit(main())
