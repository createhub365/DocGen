"""Phase 3 tenant isolation — fields, templates, mappings, generation."""
from __future__ import annotations

import io
import os
from pathlib import Path

from docx import Document

from models import FieldDefinition, GeneratedDocument, PlaceholderMapping, Template
from tests.conftest import TEST_OUTPUT_DIR, TEST_TEMPLATE_DIR


def _make_docx_bytes(*placeholder_ids: str) -> bytes:
    doc = Document()
    for pid in placeholder_ids:
        doc.add_paragraph(f"Hello {{{{{pid}}}}}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _setup_published_flow_with_field(client_a, *, slug: str, field_key: str = "cand_name"):
    dt = client_a.post(
        "/api/platform/document-types/",
        json={"name": slug, "slug": slug},
    )
    assert dt.status_code == 201, dt.text
    dt_id = dt.json()["id"]

    flow = client_a.post(f"/api/platform/{dt_id}/flow", json={})
    assert flow.status_code == 201, flow.text
    flow_id = flow.json()["id"]

    step = client_a.post(
        f"/api/platform/{flow_id}/steps",
        json={
            "step_type": "text_field",
            "order_index": 0,
            "label": "Candidate",
            "is_enabled": True,
        },
    )
    assert step.status_code == 201, step.text
    step_id = step.json()["id"]

    field = client_a.post(
        f"/api/platform/steps/{step_id}/fields",
        json={
            "field_key": field_key,
            "field_label": "Candidate Name",
            "field_type": "text",
            "is_required": True,
        },
    )
    assert field.status_code == 201, field.text

    pub = client_a.post(f"/api/platform/{flow_id}/publish")
    assert pub.status_code == 200, pub.text

    return {
        "dt_id": dt_id,
        "flow_id": flow_id,
        "step_id": step_id,
        "field_id": field.json()["id"],
        "field_key": field_key,
    }


def test_cross_org_field_definitions_blocked(dual_org_clients):
    client_a = dual_org_clients["client_a"]
    client_b = dual_org_clients["client_b"]
    db = dual_org_clients["db"]

    setup = _setup_published_flow_with_field(client_a, slug="field-guard")
    step_id = setup["step_id"]
    field_id = setup["field_id"]

    db.expire_all()
    count_before = db.query(FieldDefinition).count()

    post = client_b.post(
        f"/api/platform/steps/{step_id}/fields",
        json={
            "field_key": "hacked",
            "field_label": "Hacked",
            "field_type": "text",
        },
    )
    assert post.status_code == 404

    patch = client_b.patch(
        f"/api/platform/fields/{field_id}",
        json={"field_label": "Changed"},
    )
    assert patch.status_code == 404

    delete = client_b.delete(f"/api/platform/fields/{field_id}")
    assert delete.status_code == 404

    db.expire_all()
    assert db.query(FieldDefinition).count() == count_before
    row = db.query(FieldDefinition).filter(FieldDefinition.id == field_id).first()
    assert row is not None
    assert row.field_label == "Candidate Name"
    assert db.query(FieldDefinition).filter_by(field_key="hacked").count() == 0


def test_cross_org_template_upload_blocked(dual_org_clients):
    client_a = dual_org_clients["client_a"]
    client_b = dual_org_clients["client_b"]
    db = dual_org_clients["db"]

    setup = _setup_published_flow_with_field(client_a, slug="tmpl-guard")
    dt_id = setup["dt_id"]
    org_a_id = dual_org_clients["org_a"]["org"].id
    org_b_id = dual_org_clients["org_b"]["org"].id

    db.expire_all()
    tmpl_before = db.query(Template).count()
    files_before = list(Path(TEST_TEMPLATE_DIR).rglob("*.docx"))

    content = _make_docx_bytes("cand_name")
    resp = client_b.post(
        f"/api/platform/{dt_id}/templates",
        files={
            "file": (
                "hack.docx",
                content,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert resp.status_code == 404

    db.expire_all()
    assert db.query(Template).count() == tmpl_before
    files_after = list(Path(TEST_TEMPLATE_DIR).rglob("*.docx"))
    assert files_after == files_before
    # No org_b files written under org_a's tree either
    assert not (Path(TEST_TEMPLATE_DIR) / "orgs" / org_b_id).exists() or list(
        (Path(TEST_TEMPLATE_DIR) / "orgs" / org_b_id).rglob("*.docx")
    ) == []
    assert not list((Path(TEST_TEMPLATE_DIR) / "orgs" / org_a_id).rglob("*.docx"))


def test_template_and_generated_lists_are_org_scoped(dual_org_clients):
    client_a = dual_org_clients["client_a"]
    client_b = dual_org_clients["client_b"]

    setup_a = _setup_published_flow_with_field(client_a, slug="list-a")
    setup_b = _setup_published_flow_with_field(client_b, slug="list-b")

    content = _make_docx_bytes("cand_name")
    files = {
        "file": (
            "same_name.docx",
            content,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    up_a = client_a.post(f"/api/platform/{setup_a['dt_id']}/templates", files=files)
    up_b = client_b.post(f"/api/platform/{setup_b['dt_id']}/templates", files=files)
    assert up_a.status_code == 201, up_a.text
    assert up_b.status_code == 201, up_b.text
    a_id = up_a.json()["id"]
    b_id = up_b.json()["id"]

    list_a = client_a.get(f"/api/platform/{setup_a['dt_id']}/templates").json()
    list_b = client_b.get(f"/api/platform/{setup_b['dt_id']}/templates").json()
    assert {r["id"] for r in list_a} == {a_id}
    assert {r["id"] for r in list_b} == {b_id}
    assert b_id not in {r["id"] for r in list_a}
    assert a_id not in {r["id"] for r in list_b}

    # Complete mappings + generate for both so generated lists can be compared
    for client, setup, tmpl_id in (
        (client_a, setup_a, a_id),
        (client_b, setup_b, b_id),
    ):
        mapped = client.post(
            f"/api/platform/templates/{tmpl_id}/mappings",
            json={
                "mappings": [
                    {"placeholder_key": "cand_name", "field_key": setup["field_key"]}
                ]
            },
        )
        assert mapped.status_code == 200, mapped.text
        gen = client.post(
            f"/api/platform/{setup['dt_id']}/generate",
            json={"template_id": tmpl_id, "fields": {setup["field_key"]: "Alice"}},
        )
        assert gen.status_code == 201, gen.text

    gen_a = client_a.get("/api/platform/generated").json()
    gen_b = client_b.get("/api/platform/generated").json()
    assert len(gen_a) == 1
    assert len(gen_b) == 1
    assert gen_a[0]["org_id"] == dual_org_clients["org_a"]["org"].id
    assert gen_b[0]["org_id"] == dual_org_clients["org_b"]["org"].id
    assert gen_a[0]["id"] != gen_b[0]["id"]


def test_mapping_invalid_field_key_rejects_whole_batch(dual_org_clients):
    client_a = dual_org_clients["client_a"]
    db = dual_org_clients["db"]

    setup = _setup_published_flow_with_field(client_a, slug="map-reject")
    up = client_a.post(
        f"/api/platform/{setup['dt_id']}/templates",
        files={
            "file": (
                "t.docx",
                _make_docx_bytes("cand_name"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert up.status_code == 201, up.text
    tmpl_id = up.json()["id"]

    db.expire_all()
    before = db.query(PlaceholderMapping).count()

    resp = client_a.post(
        f"/api/platform/templates/{tmpl_id}/mappings",
        json={
            "mappings": [
                {"placeholder_key": "cand_name", "field_key": setup["field_key"]},
                {"placeholder_key": "other", "field_key": "does_not_exist"},
            ]
        },
    )
    assert resp.status_code == 400
    detail = resp.json()["detail"]
    assert "does_not_exist" in detail.get("invalid_field_keys", [])

    db.expire_all()
    assert db.query(PlaceholderMapping).count() == before


def test_generate_incomplete_mappings_blocked(dual_org_clients):
    client_a = dual_org_clients["client_a"]
    db = dual_org_clients["db"]

    setup = _setup_published_flow_with_field(client_a, slug="gen-incomplete")
    up = client_a.post(
        f"/api/platform/{setup['dt_id']}/templates",
        files={
            "file": (
                "t.docx",
                _make_docx_bytes("cand_name", "extra_ph"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert up.status_code == 201, up.text
    tmpl_id = up.json()["id"]

    # Only map one of two placeholders
    mapped = client_a.post(
        f"/api/platform/templates/{tmpl_id}/mappings",
        json={
            "mappings": [
                {"placeholder_key": "cand_name", "field_key": setup["field_key"]}
            ]
        },
    )
    assert mapped.status_code == 200
    assert mapped.json()["is_complete"] is False

    db.expire_all()
    docs_before = db.query(GeneratedDocument).count()
    files_before = list(Path(TEST_OUTPUT_DIR).rglob("*.docx"))

    resp = client_a.post(
        f"/api/platform/{setup['dt_id']}/generate",
        json={"template_id": tmpl_id, "fields": {setup["field_key"]: "Bob"}},
    )
    assert resp.status_code == 400

    db.expire_all()
    assert db.query(GeneratedDocument).count() == docs_before
    assert list(Path(TEST_OUTPUT_DIR).rglob("*.docx")) == files_before


def test_generate_without_published_flow_is_404(dual_org_clients):
    client_a = dual_org_clients["client_a"]

    dt = client_a.post(
        "/api/platform/document-types/",
        json={"name": "No Pub", "slug": "no-pub"},
    )
    assert dt.status_code == 201
    dt_id = dt.json()["id"]
    flow = client_a.post(f"/api/platform/{dt_id}/flow", json={})
    assert flow.status_code == 201
    # deliberately not published

    up = client_a.post(
        f"/api/platform/{dt_id}/templates",
        files={
            "file": (
                "t.docx",
                _make_docx_bytes("cand_name"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert up.status_code == 201

    resp = client_a.post(
        f"/api/platform/{dt_id}/generate",
        json={"fields": {"cand_name": "X"}},
    )
    assert resp.status_code == 404


def test_cross_org_generate_blocked(dual_org_clients):
    client_a = dual_org_clients["client_a"]
    client_b = dual_org_clients["client_b"]
    db = dual_org_clients["db"]

    setup = _setup_published_flow_with_field(client_a, slug="gen-xorg")
    up = client_a.post(
        f"/api/platform/{setup['dt_id']}/templates",
        files={
            "file": (
                "t.docx",
                _make_docx_bytes("cand_name"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert up.status_code == 201
    tmpl_id = up.json()["id"]
    client_a.post(
        f"/api/platform/templates/{tmpl_id}/mappings",
        json={
            "mappings": [
                {"placeholder_key": "cand_name", "field_key": setup["field_key"]}
            ]
        },
    )

    db.expire_all()
    before = db.query(GeneratedDocument).count()

    resp = client_b.post(
        f"/api/platform/{setup['dt_id']}/generate",
        json={"template_id": tmpl_id, "fields": {setup["field_key"]: "Eve"}},
    )
    assert resp.status_code == 404

    db.expire_all()
    assert db.query(GeneratedDocument).count() == before


def test_cross_org_download_is_404_without_file_bytes(dual_org_clients):
    client_a = dual_org_clients["client_a"]
    client_b = dual_org_clients["client_b"]

    setup = _setup_published_flow_with_field(client_a, slug="dl-xorg")
    up = client_a.post(
        f"/api/platform/{setup['dt_id']}/templates",
        files={
            "file": (
                "t.docx",
                _make_docx_bytes("cand_name"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    tmpl_id = up.json()["id"]
    client_a.post(
        f"/api/platform/templates/{tmpl_id}/mappings",
        json={
            "mappings": [
                {"placeholder_key": "cand_name", "field_key": setup["field_key"]}
            ]
        },
    )
    gen = client_a.post(
        f"/api/platform/{setup['dt_id']}/generate",
        json={"template_id": tmpl_id, "fields": {setup["field_key"]: "Zed"}},
    )
    assert gen.status_code == 201, gen.text
    doc_id = gen.json()["document_id"]

    # org_a can download
    ok = client_a.get(f"/api/platform/generated/{doc_id}/download")
    assert ok.status_code == 200
    assert len(ok.content) > 0
    assert ok.content[:2] == b"PK"  # docx zip magic

    denied = client_b.get(f"/api/platform/generated/{doc_id}/download")
    assert denied.status_code == 404
    # 404 body must not contain the document bytes
    assert denied.content[:2] != b"PK"
    assert b"PK" not in denied.content
    assert "Not found" in denied.text or "not found" in denied.text.lower()
