"""Bulk FieldDefinitions from template placeholders (draft flow only)."""
from __future__ import annotations

import io

from docx import Document

from models import AuditLog, FieldDefinition, FlowStep


def _docx_bytes(*placeholder_ids: str) -> bytes:
    doc = Document()
    for pid in placeholder_ids:
        doc.add_paragraph(f"Hello {{{{{pid}}}}}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_generate_fields_from_placeholders_creates_custom_fields_step(
    dual_org_clients,
):
    client = dual_org_clients["client_a"]
    db = dual_org_clients["db"]

    dt = client.post(
        "/api/platform/document-types/",
        json={"name": "Bulk Fields", "slug": "bulk-fields"},
    )
    assert dt.status_code == 201, dt.text
    dt_id = dt.json()["id"]

    flow = client.post(f"/api/platform/{dt_id}/flow", json={})
    assert flow.status_code == 201, flow.text
    flow_id = flow.json()["id"]
    # Draft with zero custom field steps / definitions

    upload = client.post(
        f"/api/platform/{dt_id}/templates",
        files={
            "file": (
                "bulk.docx",
                _docx_bytes("Company_Name", "Date", "Other_Thing"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == 201, upload.text
    template_id = upload.json()["id"]

    gen = client.post(
        f"/api/platform/templates/{template_id}/generate-fields-from-placeholders"
    )
    assert gen.status_code == 200, gen.text
    body = gen.json()
    assert body["template_id"] == template_id
    assert body["flow_config_id"] == flow_id
    assert len(body["created"]) == 3
    assert body["skipped_placeholders"] == []
    assert body["possible_duplicates"] == []

    by_key = {item["field_key"]: item["field_label"] for item in body["created"]}
    assert by_key["company_name"] == "Company Name"
    assert by_key["date"] == "Date"
    assert by_key["other_thing"] == "Other Thing"

    db.expire_all()
    step = (
        db.query(FlowStep)
        .filter(
            FlowStep.flow_config_id == flow_id,
            FlowStep.step_type == "custom_fields",
        )
        .one()
    )
    assert step.label == "Generated fields"
    assert body["flow_step_id"] == step.id

    fields = (
        db.query(FieldDefinition)
        .filter(FieldDefinition.flow_step_id == step.id)
        .order_by(FieldDefinition.id.asc())
        .all()
    )
    assert len(fields) == 3
    for fd in fields:
        assert fd.field_type == "text"
        assert fd.is_required is True

    audits = (
        db.query(AuditLog)
        .filter(AuditLog.action == "fields.bulk_generated_from_template")
        .all()
    )
    assert len(audits) == 1
    assert audits[0].target_type == "Template"
    assert audits[0].target_id == str(template_id)
    assert audits[0].metadata_json["created_count"] == 3

    # Idempotency: second call creates nothing
    again = client.post(
        f"/api/platform/templates/{template_id}/generate-fields-from-placeholders"
    )
    assert again.status_code == 200, again.text
    again_body = again.json()
    assert again_body["created"] == []
    assert set(again_body["skipped_placeholders"]) == {
        "Company_Name",
        "Date",
        "Other_Thing",
    }
    assert again_body["possible_duplicates"] == []

    db.expire_all()
    assert (
        db.query(FieldDefinition)
        .filter(FieldDefinition.flow_step_id == step.id)
        .count()
        == 3
    )
    assert (
        db.query(AuditLog)
        .filter(AuditLog.action == "fields.bulk_generated_from_template")
        .count()
        == 2
    )


def test_generate_fields_flags_fuzzy_duplicates_for_review(dual_org_clients):
    """Near-duplicate placeholders are held for review, not auto-created."""
    client = dual_org_clients["client_a"]
    db = dual_org_clients["db"]

    dt = client.post(
        "/api/platform/document-types/",
        json={"name": "Fuzzy Dup Gen", "slug": "fuzzy-dup-gen"},
    )
    assert dt.status_code == 201, dt.text
    dt_id = dt.json()["id"]
    flow = client.post(f"/api/platform/{dt_id}/flow", json={})
    assert flow.status_code == 201
    flow_id = flow.json()["id"]

    # Seed existing field "position" on a draft custom_fields step
    step_resp = client.post(
        f"/api/platform/{flow_id}/steps",
        json={
            "step_type": "custom_fields",
            "order_index": 0,
            "is_enabled": True,
            "label": "Details",
        },
    )
    assert step_resp.status_code == 201, step_resp.text
    step_id = step_resp.json()["id"]
    field_resp = client.post(
        f"/api/platform/steps/{step_id}/fields",
        json={
            "field_key": "position",
            "field_label": "Position",
            "field_type": "text",
            "is_required": False,
        },
    )
    assert field_resp.status_code == 201, field_resp.text

    upload = client.post(
        f"/api/platform/{dt_id}/templates",
        files={
            "file": (
                "offer.docx",
                _docx_bytes(
                    "position_title",
                    "employee_email",
                    "department_code",
                ),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == 201, upload.text
    template_id = upload.json()["id"]

    gen = client.post(
        f"/api/platform/templates/{template_id}/generate-fields-from-placeholders"
    )
    assert gen.status_code == 200, gen.text
    body = gen.json()

    created_keys = {item["field_key"] for item in body["created"]}
    assert "position_title" not in created_keys
    assert "employee_email" in created_keys
    assert "department_code" in created_keys
    assert "position_title" not in body["skipped_placeholders"]

    possibles = body["possible_duplicates"]
    assert len(possibles) == 1
    assert possibles[0]["placeholder"] == "position_title"
    assert possibles[0]["proposed_field_key"] == "position_title"
    assert "position" in possibles[0]["similar_field_keys"]

    db.expire_all()
    keys_before_confirm = {
        row.field_key
        for row in db.query(FieldDefinition)
        .filter(FieldDefinition.flow_step_id == step_id)
        .all()
    }
    assert "position_title" not in keys_before_confirm
    assert "employee_email" in keys_before_confirm

    # Explicit confirm creates the held field
    confirm = client.post(
        f"/api/platform/templates/{template_id}/generate-fields-from-placeholders",
        json={"create_placeholders": ["position_title"]},
    )
    assert confirm.status_code == 200, confirm.text
    confirm_body = confirm.json()
    assert any(c["field_key"] == "position_title" for c in confirm_body["created"])
    assert confirm_body["possible_duplicates"] == []

    db.expire_all()
    keys_after = {
        row.field_key
        for row in db.query(FieldDefinition)
        .filter(FieldDefinition.flow_step_id == step_id)
        .all()
    }
    assert "position_title" in keys_after


def test_generate_fields_skips_ref_number_barcode_placeholder(dual_org_clients):
    """Bulk generate must not create a separate barcode FieldDefinition."""
    client = dual_org_clients["client_a"]
    db = dual_org_clients["db"]

    dt = client.post(
        "/api/platform/document-types/",
        json={"name": "Barcode Skip", "slug": "barcode-skip"},
    )
    assert dt.status_code == 201, dt.text
    dt_id = dt.json()["id"]
    flow = client.post(f"/api/platform/{dt_id}/flow", json={})
    assert flow.status_code == 201
    flow_id = flow.json()["id"]

    upload = client.post(
        f"/api/platform/{dt_id}/templates",
        files={
            "file": (
                "ref.docx",
                _docx_bytes("ref_number", "ref_number_barcode", "cand_name"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == 201, upload.text
    template_id = upload.json()["id"]

    gen = client.post(
        f"/api/platform/templates/{template_id}/generate-fields-from-placeholders"
    )
    assert gen.status_code == 200, gen.text
    body = gen.json()
    created_keys = {item["field_key"] for item in body["created"]}
    assert "ref_number_barcode" not in created_keys
    assert "ref_number_barcode" in body["skipped_placeholders"]
    assert "ref_number" in created_keys
    assert "cand_name" in created_keys

    db.expire_all()
    keys = {
        row.field_key
        for row in db.query(FieldDefinition)
        .join(FlowStep, FlowStep.id == FieldDefinition.flow_step_id)
        .filter(FlowStep.flow_config_id == flow_id)
        .all()
    }
    assert "ref_number_barcode" not in keys
    assert "ref_number" in keys


def test_generate_fields_requires_draft_flow(dual_org_clients):
    client = dual_org_clients["client_a"]

    dt = client.post(
        "/api/platform/document-types/",
        json={"name": "No Draft Gen", "slug": "no-draft-gen"},
    )
    assert dt.status_code == 201
    dt_id = dt.json()["id"]

    flow = client.post(f"/api/platform/{dt_id}/flow", json={})
    assert flow.status_code == 201
    flow_id = flow.json()["id"]
    pub = client.post(f"/api/platform/{flow_id}/publish")
    assert pub.status_code == 200

    upload = client.post(
        f"/api/platform/{dt_id}/templates",
        files={
            "file": (
                "x.docx",
                _docx_bytes("only_one"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == 201
    template_id = upload.json()["id"]

    gen = client.post(
        f"/api/platform/templates/{template_id}/generate-fields-from-placeholders"
    )
    assert gen.status_code == 404
    assert "draft" in gen.json()["detail"].lower()


def test_generate_fields_cross_org_is_404(dual_org_clients):
    client_a = dual_org_clients["client_a"]
    client_b = dual_org_clients["client_b"]

    dt = client_a.post(
        "/api/platform/document-types/",
        json={"name": "XOrg Gen", "slug": "xorg-gen"},
    )
    assert dt.status_code == 201
    dt_id = dt.json()["id"]
    assert client_a.post(f"/api/platform/{dt_id}/flow", json={}).status_code == 201
    upload = client_a.post(
        f"/api/platform/{dt_id}/templates",
        files={
            "file": (
                "x.docx",
                _docx_bytes("a_key"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == 201
    template_id = upload.json()["id"]

    blocked = client_b.post(
        f"/api/platform/templates/{template_id}/generate-fields-from-placeholders"
    )
    assert blocked.status_code == 404
